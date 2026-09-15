from pathlib import Path
from io import BytesIO
import csv

import fitz
from docx import Document
from PIL import Image
import pytesseract


def extract_pdf_text(file_path: Path) -> str:
    text_parts = []

    with fitz.open(file_path) as pdf:
        for page in pdf:
            text = page.get_text()

            if text and text.strip():
                text_parts.append(text.strip())

    return "\n".join(text_parts)


def extract_docx_text(file_path: Path) -> str:
    document = Document(file_path)

    text_parts = []

    # Normal paragraphs
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()

        if text:
            text_parts.append(text)

    # Tables
    for table in document.tables:
        for row in table.rows:
            row_values = []

            for cell in row.cells:
                cell_text = cell.text.strip()

                if cell_text:
                    row_values.append(cell_text)

            if row_values:
                text_parts.append(
                    " | ".join(row_values)
                )

    normal_text = "\n".join(text_parts).strip()

    if normal_text:
        return normal_text

    print(
        "No normal DOCX text found. "
        "Trying OCR on embedded images..."
    )

    # OCR fallback for scanned/image-based Word docs
    ocr_parts = []

    for relationship in document.part.rels.values():
        if "image" not in relationship.reltype:
            continue

        try:
            image_blob = relationship.target_part.blob

            image = Image.open(
                BytesIO(image_blob)
            )

            image_text = pytesseract.image_to_string(
                image
            ).strip()

            if image_text:
                ocr_parts.append(image_text)

        except Exception as exc:
            print(
                "DOCX IMAGE OCR ERROR:",
                str(exc)
            )

    return "\n".join(ocr_parts)


def extract_txt_text(file_path: Path) -> str:
    return file_path.read_text(
        encoding="utf-8",
        errors="ignore",
    )


def extract_csv_text(file_path: Path) -> str:
    rows = []

    with open(
        file_path,
        "r",
        encoding="utf-8-sig",
        errors="ignore",
        newline="",
    ) as csv_file:

        reader = csv.reader(csv_file)

        for row in reader:
            clean_row = [
                cell.strip()
                for cell in row
                if cell.strip()
            ]

            if clean_row:
                rows.append(
                    " | ".join(clean_row)
                )

    return "\n".join(rows)


def extract_text_from_file(
    file_path: Path
) -> str:

    extension = file_path.suffix.lower()

    print(
        "DOCUMENT SERVICE EXTENSION:",
        extension
    )

    if extension == ".pdf":
        return extract_pdf_text(file_path)

    if extension == ".docx":
        return extract_docx_text(file_path)

    if extension == ".txt":
        return extract_txt_text(file_path)

    if extension == ".csv":
        return extract_csv_text(file_path)

    raise ValueError(
        f"Unsupported file type: {extension}"
    )